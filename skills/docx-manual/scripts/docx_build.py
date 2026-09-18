#!/usr/bin/env python3
"""Build a .docx manual from Markdown + screenshots, using an existing .docx for style.

The template document supplies the look: its heading/body/caption paragraphs,
horizontal rule, page-break paragraph and first table are cloned, and its
``sectPr`` (page size, margins, headers, footers) is kept.  The template's body
text is cleared before the new content is written.

Supported Markdown subset::

    # / ## / ###            headings
    text with **bold** and *italic*
    - item   * item         bullet list
    1. item                 numbered list
    | a | b |               pipe table (first row = header)
    ![caption](shot.png)    figure with an auto-numbered caption
    ![a](x.png) ![b](y.png) two pictures side by side under one caption
    > quoted text           body paragraph in italics
    ---                     horizontal rule
    <!-- pagebreak -->      page break

Relative image paths resolve against ``--images-dir`` first, then the directory
of the Markdown file.  If any image is missing, nothing is written and every
missing path is listed.

Example::

    python docx_build.py manual.md --template old-manual.docx --out new.docx \\
        --images-dir shots --width-cm 15 --break-before-h1
"""

from __future__ import annotations

import argparse
import copy
import json
import os
import re
import sys

import docx

sys.path.insert(0, __file__.rsplit('/', 1)[0])
from docxkit import (  # noqa: E402
    FIGURE_FORMAT, FIGURE_PATTERN, TemplateSet, add_figure, append_elements,
    clear_body, clone_hr, clone_pagebreak, clone_para_runs, clone_heading,
    clone_table, human_size, prune_unused_images,
)

RE_HEADING = re.compile(r'^(#{1,3})\s+(.*)$')
RE_BULLET = re.compile(r'^[-*]\s+(.*)$')
RE_NUMBER = re.compile(r'^\d+[.)]\s+(.*)$')
RE_FIGURE = re.compile(r'!\[([^\]]*)\]\(([^)]+)\)')
RE_FIGURE_LINE = re.compile(r'^\s*(!\[[^\]]*\]\([^)]+\)\s*){1,2}$')
RE_INLINE = re.compile(r'(\*\*[^*]+\*\*|\*[^*]+\*)')
RE_HR = re.compile(r'^-{3,}$')
RE_PAGEBREAK = re.compile(r'^<!--\s*pagebreak\s*-->$', re.I)


# --------------------------------------------------------------------------
# Markdown → blocks
# --------------------------------------------------------------------------

def inline_parts(text):
    """``'a **b** c'`` → ``[('a ', {}), ('b', {'bold': True}), (' c', {})]``."""
    parts = []
    for chunk in RE_INLINE.split(text):
        if not chunk:
            continue
        if chunk.startswith('**') and chunk.endswith('**'):
            parts.append((chunk[2:-2], {'bold': True}))
        elif chunk.startswith('*') and chunk.endswith('*') and len(chunk) > 2:
            parts.append((chunk[1:-1], {'italic': True}))
        else:
            parts.append((chunk, {}))
    return parts or [('', {})]


def parse(md_text):
    """Parse the Markdown subset into a flat list of block tuples."""
    lines = md_text.replace('\r\n', '\n').split('\n')
    blocks = []
    buf = []          # pending plain-paragraph lines
    table = []        # pending table rows
    i = 0

    def flush_para():
        if buf:
            blocks.append(('para', ' '.join(buf).strip()))
            del buf[:]

    def flush_table():
        if table:
            blocks.append(('table', list(table)))
            del table[:]

    while i < len(lines):
        line = lines[i].rstrip()
        stripped = line.strip()
        i += 1

        if stripped.startswith('|') and stripped.endswith('|'):
            flush_para()
            cells = [c.strip() for c in stripped.strip('|').split('|')]
            if not all(re.fullmatch(r':?-{2,}:?', c) for c in cells):
                table.append(cells)
            continue
        flush_table()

        if not stripped:
            flush_para()
            continue
        if RE_PAGEBREAK.match(stripped):
            flush_para()
            blocks.append(('pagebreak', None))
            continue
        if RE_HR.match(stripped):
            flush_para()
            blocks.append(('hr', None))
            continue
        m = RE_HEADING.match(stripped)
        if m:
            flush_para()
            blocks.append(('heading', (len(m.group(1)), m.group(2).strip())))
            continue
        if RE_FIGURE_LINE.match(stripped):
            flush_para()
            found = RE_FIGURE.findall(stripped)
            alts = [a for a, _ in found]
            paths = [p.strip() for _, p in found]
            blocks.append(('figure', (alts, paths)))
            continue
        if stripped.startswith('>'):
            flush_para()
            blocks.append(('quote', stripped.lstrip('>').strip()))
            continue
        m = RE_BULLET.match(stripped)
        if m:
            flush_para()
            blocks.append(('bullet', m.group(1).strip()))
            continue
        m = RE_NUMBER.match(stripped)
        if m:
            flush_para()
            blocks.append(('number', m.group(1).strip()))
            continue
        buf.append(stripped)

    flush_para()
    flush_table()
    return blocks


# --------------------------------------------------------------------------
# blocks → docx
# --------------------------------------------------------------------------

class Builder:
    """Renders parsed blocks into `doc` using cloned elements from `templates`."""

    def __init__(self, doc, templates, width_cm=None, figure_format=FIGURE_FORMAT,
                 break_before_h1=False):
        self.doc = doc
        self.t = templates
        self.width_cm = width_cm
        self.figure_format = figure_format
        self.break_before_h1 = break_before_h1
        self.figure_no = 0
        self.headings = 0
        self._style_ids = {s.style_id for s in doc.styles}
        self._seen_h1 = False

    # -- helpers -----------------------------------------------------------
    def _add(self, el):
        append_elements(self.doc, [el])

    def _styled(self, style_id, parts):
        """Paragraph in a named style (used for list styles the template defines)."""
        p = self.doc.add_paragraph(style=style_id)
        for text, fmt in parts:
            r = p.add_run(text)
            r.bold = bool(fmt.get('bold'))
            r.italic = bool(fmt.get('italic'))
        return p

    # -- blocks ------------------------------------------------------------
    def heading(self, level, text):
        self.headings += 1
        if level == 1 and self.break_before_h1 and self._seen_h1:
            self.pagebreak()
        if level == 1:
            self._seen_h1 = True
        tpl = getattr(self.t, 'h%d' % level)
        if tpl is not None:
            self._add(clone_heading(tpl, text))
            return
        style = 'Heading %d' % level
        if style in self._style_ids or style.replace(' ', '') in self._style_ids:
            self.doc.add_paragraph(text, style=style)
        else:  # last resort: bold body text
            self.para('**%s**' % text)

    def para(self, text, parts=None):
        parts = parts if parts is not None else inline_parts(text)
        if self.t.body is not None:
            self._add(clone_para_runs(self.t.body, parts))
        else:
            self._styled('Normal', parts)

    def quote(self, text):
        parts = [(t, dict(f, italic=True)) for t, f in inline_parts(text)]
        self.para(None, parts)

    def listitem(self, kind, text):
        parts = inline_parts(text)
        style = 'List Bullet' if kind == 'bullet' else 'List Number'
        if style in self._style_ids or style.replace(' ', '') in self._style_ids:
            self._styled(style, parts)
            return
        tpl = getattr(self.t, kind)
        if tpl is not None:
            self._add(clone_para_runs(tpl, parts))
            return
        # no list style and no numbered paragraph to clone: fake a bullet
        bullet = '• ' if kind == 'bullet' else '– '
        self.para(None, [(bullet, {})] + parts)

    def table(self, rows):
        if self.t.table is not None and rows:
            width = len(self.t.table.findall(
                '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tblGrid'
                '/{http://schemas.openxmlformats.org/wordprocessingml/2006/main}gridCol'))
            if width == len(rows[0]):
                self._add(clone_table(self.t.table, rows))
                return
        t = self.doc.add_table(rows=len(rows), cols=len(rows[0]))
        if 'TableGrid' in self._style_ids or 'Table Grid' in self._style_ids:
            t.style = 'Table Grid'
        for r, vals in enumerate(rows):
            for c, v in enumerate(vals):
                t.cell(r, c).text = v

    def hr(self):
        # clone_hr falls back to a bottom-border paragraph when the template
        # document has no rule of its own.
        self._add(clone_hr(self.t.hr))

    def pagebreak(self):
        self._add(clone_pagebreak(self.t.pagebreak))

    def figure(self, alts, paths):
        caption = self._caption(alts)
        add_figure(self.doc, paths[0], caption, width_cm=self.width_cm,
                   caption_template=self.t.caption,
                   side_by_side_paths=paths if len(paths) == 2 else None)

    def _caption(self, alts):
        """Number the caption unless the alt text already carries a figure number."""
        alts = [a for a in alts if a]
        if len(alts) == 2:
            alt = '%s｜%s' % (alts[0], alts[1])
        else:
            alt = alts[0] if alts else ''
        self.figure_no += 1
        if re.match(FIGURE_PATTERN, alt):
            return alt
        return self.figure_format.format(n=self.figure_no) + alt

    def run(self, blocks):
        for kind, payload in blocks:
            if kind == 'heading':
                self.heading(*payload)
            elif kind == 'para':
                self.para(payload)
            elif kind == 'quote':
                self.quote(payload)
            elif kind in ('bullet', 'number'):
                self.listitem(kind, payload)
            elif kind == 'table':
                self.table(payload)
            elif kind == 'hr':
                self.hr()
            elif kind == 'pagebreak':
                self.pagebreak()
            elif kind == 'figure':
                self.figure(*payload)
            else:  # pragma: no cover
                raise AssertionError('unknown block %r' % kind)


# --------------------------------------------------------------------------
# CLI
# --------------------------------------------------------------------------

def resolve_images(blocks, images_dir, md_dir):
    """Rewrite figure paths in place; return the list of unresolvable ones."""
    missing = []
    for idx, (kind, payload) in enumerate(blocks):
        if kind != 'figure':
            continue
        alts, paths = payload
        fixed = []
        for p in paths:
            if os.path.isabs(p) and os.path.exists(p):
                fixed.append(p)
                continue
            for base in [d for d in (images_dir, md_dir) if d]:
                cand = os.path.join(base, p)
                if os.path.exists(cand):
                    fixed.append(cand)
                    break
            else:
                if os.path.exists(p):
                    fixed.append(p)
                else:
                    missing.append(p)
                    fixed.append(p)
        blocks[idx] = (kind, (alts, fixed))
    return missing


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('source', help='Markdown source file')
    ap.add_argument('--template', required=True, help='.docx whose styles are reused')
    ap.add_argument('--out', required=True, help='output .docx')
    ap.add_argument('--images-dir', help='directory searched first for relative images')
    ap.add_argument('--figure-format', default=FIGURE_FORMAT,
                    help="caption prefix format, default %r" % FIGURE_FORMAT)
    ap.add_argument('--width-cm', type=float, help='displayed picture width in cm')
    ap.add_argument('--break-before-h1', action='store_true',
                    help='start every H1 after the first on a new page')
    ap.add_argument('--style-map', metavar='JSON',
                    help='JSON {"h1": "Heading1", ...} overriding template discovery')
    ap.add_argument('--title', help='core-properties title of the output')
    args = ap.parse_args(argv)

    with open(args.source, encoding='utf-8') as fh:
        blocks = parse(fh.read())

    md_dir = os.path.dirname(os.path.abspath(args.source))
    missing = resolve_images(blocks, args.images_dir, md_dir)
    if missing:
        raise SystemExit('missing image(s), nothing written:\n  ' + '\n  '.join(missing))

    doc = docx.Document(args.template)
    style_map = json.load(open(args.style_map, encoding='utf-8')) if args.style_map else None
    templates = TemplateSet.from_doc(doc, style_overrides=style_map)
    # Detach the templates before clearing the body, so cloning still works.
    for field in TemplateSet.FIELDS:
        el = getattr(templates, field)
        setattr(templates, field, copy.deepcopy(el) if el is not None else None)
    clear_body(doc)

    b = Builder(doc, templates, width_cm=args.width_cm,
                figure_format=args.figure_format,
                break_before_h1=args.break_before_h1)
    b.run(blocks)

    # The template's own screenshots are still in the package after
    # clear_body(): a 20 MB client manual used as a template would otherwise
    # produce a 20 MB two-figure document.
    dropped = prune_unused_images(doc)

    if args.title:
        doc.core_properties.title = args.title
    doc.save(args.out)
    print('-- %s: %d heading(s), %d figure(s), %d block(s), '
          '%d template image part(s) pruned, %s'
          % (args.out, b.headings, b.figure_no, len(blocks), dropped,
             human_size(args.out)))
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
